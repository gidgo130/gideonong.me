"""Read a CV master (.docx) into paragraphs with resolved formatting.

Read-only. For every paragraph: its text, a kind (name | title | contact |
heading | entry | bullet | plain), its runs with bold / italic / size resolved
the way Word resolves them (run → character style → paragraph style chain →
document defaults; the paragraph-mark rPr never applies to runs), the right
tab stop that positions a date, and the left indent (direct → numbering level
→ style chain). The section gives page width and margins so the fit rule uses
the document's own geometry, never hard-coded numbers.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn

TWIP = 20.0  # twips per point


@dataclass
class Run:
    text: str
    bold: bool
    italic: bool
    underline: bool
    size_pt: float
    font: str


@dataclass
class Para:
    index: int
    text: str
    kind: str
    runs: list[Run] = field(default_factory=list)
    centered: bool = False
    is_bullet: bool = False
    tab_right_pt: Optional[float] = None
    left_indent_pt: float = 0.0
    right_indent_pt: float = 0.0

    @property
    def is_entry(self) -> bool:
        return self.kind == "entry"

    def to_json(self) -> dict:
        return {"index": self.index, "text": self.text, "kind": self.kind}


@dataclass
class DocInfo:
    path: str
    page_width_pt: float
    page_height_pt: float
    left_margin_pt: float
    right_margin_pt: float
    body_size_pt: float
    default_font: str
    paragraphs: list[Para]

    @property
    def text_width_pt(self) -> float:
        return self.page_width_pt - self.left_margin_pt - self.right_margin_pt

    def entries(self) -> list[Para]:
        return [p for p in self.paragraphs if p.kind == "entry"]

    def nonempty(self) -> list[Para]:
        return [p for p in self.paragraphs if p.text.strip()]


# ---------------------------------------------------------------- resolvers


def _doc_defaults(doc) -> tuple[Optional[float], str]:
    """(size_pt, font) from styles.xml docDefaults."""
    el = doc.styles.element
    size, font = None, ""
    for sz in el.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:sz"):
        try:
            size = int(sz.get(qn("w:val"))) / 2.0
        except (TypeError, ValueError):
            pass
    for rf in el.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts"):
        font = rf.get(qn("w:ascii")) or rf.get(qn("w:hAnsi")) or ""
    return size, font


def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _style_attr(style, getter):
    for s in _style_chain(style):
        try:
            v = getter(s)
        except AttributeError:
            v = None
        if v is not None:
            return v
    return None


def _run_font(run, para_style, defaults) -> tuple[bool, bool, bool, float, str]:
    def pick(getter):
        v = getter(run.font)
        if v is None and run.style is not None:
            v = _style_attr(run.style, lambda s: getter(s.font))
        if v is None:
            v = _style_attr(para_style, lambda s: getter(s.font))
        return v

    bold = pick(lambda f: f.bold)
    italic = pick(lambda f: f.italic)
    underline = pick(lambda f: f.underline)
    size = pick(lambda f: f.size)
    name = pick(lambda f: f.name)
    size_pt = size.pt if size is not None else (defaults[0] or 11.0)
    return bool(bold), bool(italic), bool(underline), float(size_pt), name or defaults[1] or ""


def _numbering_indent(doc, num_id: str, ilvl: str) -> Optional[float]:
    """Left indent (pt) from numbering.xml for numId/ilvl, or None."""
    try:
        numbering = doc.part.numbering_part.element
    except (AttributeError, KeyError, NotImplementedError):
        return None
    num = next((n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id), None)
    if num is None:
        return None

    def lvl_left(lvl) -> Optional[float]:
        # the schema puts w:ind under w:lvl/w:pPr; tolerate a direct child too
        pPr = lvl.find(qn("w:pPr"))
        ind = pPr.find(qn("w:ind")) if pPr is not None else None
        if ind is None:
            ind = lvl.find(qn("w:ind"))
        if ind is not None and ind.get(qn("w:left")) is not None:
            return int(ind.get(qn("w:left"))) / TWIP
        return None

    # a level override on the num itself wins
    for ov in num.findall(qn("w:lvlOverride")):
        if ov.get(qn("w:ilvl")) == ilvl:
            lvl = ov.find(qn("w:lvl"))
            left = lvl_left(lvl) if lvl is not None else None
            if left is not None:
                return left
    abs_el = num.find(qn("w:abstractNumId"))
    abs_id = abs_el.get(qn("w:val")) if abs_el is not None else None
    for an in numbering.findall(qn("w:abstractNum")):
        if an.get(qn("w:abstractNumId")) != abs_id:
            continue
        for lvl in an.findall(qn("w:lvl")):
            if lvl.get(qn("w:ilvl")) == ilvl:
                return lvl_left(lvl)
    return None


def _indents(doc, p) -> tuple[float, float, bool]:
    """(left_pt, right_pt, is_bullet) for a paragraph."""
    pf = p.paragraph_format
    left = pf.left_indent.pt if pf.left_indent is not None else None
    right = pf.right_indent.pt if pf.right_indent is not None else None
    num = p._p.pPr.find(qn("w:numPr")) if p._p.pPr is not None else None
    if num is None:  # a list style ("List Bullet") carries the numPr on the style
        for s in _style_chain(p.style):
            spPr = getattr(s.element, "pPr", None)
            snum = spPr.find(qn("w:numPr")) if spPr is not None else None
            if snum is not None:
                num = snum
                break
    is_bullet = num is not None
    if left is None and num is not None:
        nid, lvl = num.find(qn("w:numId")), num.find(qn("w:ilvl"))
        if nid is not None:
            left = _numbering_indent(doc, nid.get(qn("w:val")), lvl.get(qn("w:val")) if lvl is not None else "0")
    if left is None:
        left = _style_attr(p.style, lambda s: s.paragraph_format.left_indent)
        left = left.pt if left is not None else 0.0
    if right is None:
        right = _style_attr(p.style, lambda s: s.paragraph_format.right_indent)
        right = right.pt if right is not None else 0.0
    return float(left), float(right), is_bullet


def _right_tab(p) -> Optional[float]:
    stops = list(p.paragraph_format.tab_stops)
    if not stops:
        for s in _style_chain(p.style):
            stops = list(s.paragraph_format.tab_stops)
            if stops:
                break
    rights = [t.position.pt for t in stops if t.alignment == WD_TAB_ALIGNMENT.RIGHT]
    return max(rights) if rights else None


def _centered(p) -> bool:
    a = p.alignment
    if a is None:
        a = _style_attr(p.style, lambda s: s.paragraph_format.alignment)
    return a == WD_ALIGN_PARAGRAPH.CENTER


# ---------------------------------------------------------------- reading


def read(path: Path) -> DocInfo:
    doc = Document(str(path))
    defaults = _doc_defaults(doc)
    sec = doc.sections[0]
    paras: list[Para] = []
    for i, p in enumerate(doc.paragraphs):
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            b, it, u, sz, name = _run_font(r, p.style, defaults)
            runs.append(Run(r.text, b, it, u, sz, name))
        left, right, is_bullet = _indents(doc, p)
        paras.append(
            Para(
                index=i,
                text=p.text,
                kind="plain",
                runs=runs,
                centered=_centered(p),
                is_bullet=is_bullet,
                tab_right_pt=_right_tab(p),
                left_indent_pt=left,
                right_indent_pt=right,
            )
        )
    sizes = Counter(round(r.size_pt, 1) for p in paras for r in p.runs if not p.is_bullet or True)
    body = sizes.most_common(1)[0][0] if sizes else (defaults[0] or 11.0)
    info = DocInfo(
        path=str(path),
        page_width_pt=sec.page_width.pt,
        page_height_pt=sec.page_height.pt,
        left_margin_pt=sec.left_margin.pt,
        right_margin_pt=sec.right_margin.pt,
        body_size_pt=float(body),
        default_font=defaults[1],
        paragraphs=paras,
    )
    _classify(info)
    return info


def _classify(info: DocInfo) -> None:
    body = info.body_size_pt
    seen_name = False
    for p in info.paragraphs:
        text = p.text
        if not text.strip():
            p.kind = "empty"
            continue
        if p.is_bullet:
            p.kind = "bullet"
            continue
        if "\t" in text and text.rsplit("\t", 1)[1].strip():
            p.kind = "entry"
            continue
        if p.centered and not seen_name:
            p.kind = "name"
            seen_name = True
            continue
        if p.centered and p.index < 4:
            p.kind = "contact" if ("@" in text or "|" in text) else "title"
            continue
        nonempty = [r for r in p.runs if r.text.strip()]
        if nonempty and all(r.bold for r in nonempty) and max(r.size_pt for r in nonempty) > body:
            p.kind = "heading"
            continue
        p.kind = "plain"


def kind_counts(info: DocInfo) -> dict:
    c = Counter(p.kind for p in info.paragraphs if p.kind != "empty")
    return dict(sorted(c.items()))
