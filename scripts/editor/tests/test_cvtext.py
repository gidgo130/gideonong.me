"""Phase 4a: master import, pairing, sharing across variants, and the proof of
losslessness. Synthetic masters carry the same shapes as Gideon's (name / title
/ contact, headings, entry lines with a right tab, bullets, split runs with
rsid attributes and spell-check markers). The real masters are used when present."""

import copy
import json
import tempfile
import unittest
from pathlib import Path

from _helpers import REPO_ROOT, TempRepo
from core.backups import Backups
from core.cv import importer, masters, renderer
from core.cv.textservice import CvTextService

MASTERS_PRESENT = all(masters.docx_path(REPO_ROOT, m).is_file() for m in masters.MASTERS)
W = importer.W_NS


def make_master(path: Path, lang="en", title=True, entries=None, extra_bullet=True, split_runs=True):
    """A CV-shaped master: name, [title], contact, Education heading + entry + line, Experience + entries with bullets."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.32)
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(11)

    def para(text="", size=None, bold=False, italic=False, center=False, style=None):
        p = doc.add_paragraph(style=style)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text:
            r = p.add_run(text)
            if bold:
                r.bold = True
            if italic:
                r.italic = True
            if size:
                r.font.size = Pt(size)
        return p

    def entry(role, org, date):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.86), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(role)
        r.bold = True
        if org:
            p.add_run(" · " + org)
        p.add_run("\t")
        r = p.add_run(date.replace(" ", "\xa0"))
        r.italic = r.underline = True
        return p

    para("Test Person", size=18, bold=True, center=True)
    if title:
        para("Engineer · Tester" if lang == "en" else "Ingeniero · Probador", italic=True, center=True)
    para("test@example.com  |  example.com", size=10, center=True)
    para("Education" if lang == "en" else "Formación", size=13, bold=True)
    entry("Test University" if lang == "en" else "Universidad de Prueba", "", "In progress" if lang == "en" else "En curso")
    para("Program line" if lang == "en" else "Línea del programa")
    para("Experience" if lang == "en" else "Experiencia", size=13, bold=True)
    for role, org, date, bullets in (entries or [("Engineer", "Org, Town", "Summer 2026", ["Did a thing.", "Did another."])]):
        entry(role, org, date)
        for b in bullets:
            para(b, style="List Bullet")
    if extra_bullet:
        pass
    doc.save(str(path))
    if split_runs:
        # split the first bullet's run in two (rsid-only difference) and add a spell-check marker between them
        import zipfile
        import re

        z = zipfile.ZipFile(path)
        xml = z.read("word/document.xml").decode("utf-8")
        parts = {n: z.read(n) for n in z.namelist()}
        z.close()
        target = "Did a thing." if lang == "en" else None
        if target and target in xml:
            xml = xml.replace(f"<w:t>{target}</w:t></w:r>", '<w:t>Did a</w:t></w:r><w:proofErr w:type="spellStart"/><w:r w:rsidRPr="00AB12CD"><w:t xml:space="preserve"> thing.</w:t></w:r>', 1)
        parts["word/document.xml"] = xml.encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
            for n, data in parts.items():
                out.writestr(n, data)
    return path


class ReadAndStructureTests(unittest.TestCase):
    def test_read_merges_split_runs_and_classifies(self):
        with TempRepo() as t:
            p = make_master(t.root / "m.docx")
            d = importer.read_master(p)
            kinds = [x.kind for x in d.paragraphs]
            self.assertEqual(kinds, ["name", "title", "contact", "heading", "entry", "plain", "heading", "entry", "bullet", "bullet"])
            first_bullet = d.paragraphs[8]
            self.assertEqual([s.text for s in first_bullet.spans], ["Did a thing."])  # rsid split + proofErr merged away
            e = d.paragraphs[7]
            self.assertEqual([(importer.is_bold(s.key), s.text) for s in e.spans], [(True, "Engineer"), (False, " · Org, Town\t"), (False, "Summer\xa02026")])
            self.assertEqual(importer.split_entry(e.spans), {"role": "Engineer", "org": "Org, Town", "date": "Summer 2026", "sep": " · ", "notes": []})
            self.assertEqual(importer.split_entry(d.paragraphs[4].spans)["org"], "")
            nodes = importer.structure(d)
            self.assertEqual([n.kind for n in nodes], ["name", "title", "contact", "heading", "entry", "child-line", "heading", "entry", "child-bullet", "child-bullet"])

    def test_unsupported_constructs_are_refused(self):
        with TempRepo() as t:
            p = make_master(t.root / "m.docx", split_runs=False)
            from docx import Document

            doc = Document(str(p))
            doc.add_table(rows=1, cols=1)
            doc.save(str(p))
            with self.assertRaises(importer.ImportError_):
                importer.read_master(p)


class MergeAndProofTests(unittest.TestCase):
    def _docs(self, t):
        en = make_master(t.root / "Gideon Ong CV Full EN.docx", "en", entries=[("Engineer", "Org, Town", "Summer 2026", ["Did a thing.", "Did another."]), ("Tutor", "School", "Fall 2025", ["Taught."])])
        es = make_master(t.root / "Gideon Ong CV Completo ES.docx", "es", entries=[("Ingeniero", "Org, Pueblo", "Verano 2026", ["Hice algo.", "Hice otra cosa."]), ("Tutor", "Escuela", "Otoño 2025", ["Enseñé."])])
        pen = make_master(t.root / "Gideon Ong CV Professional EN.docx", "en", entries=[("Engineer", "Org, Town", "Summer 2026", ["Did a thing."])])
        pes = make_master(t.root / "Gideon Ong CV Profesional ES.docx", "es", entries=[("Ingeniero", "Org, Pueblo", "Verano 2026", ["Hice algo."])], split_runs=False)
        return {"full": {"en": importer.read_master(en), "es": importer.read_master(es)}, "professional": {"en": importer.read_master(pen), "es": importer.read_master(pes)}}

    def test_merge_shares_items_and_keeps_per_variant_order(self):
        with TempRepo() as t:
            content, slots, report = importer.merge(self._docs(t))
            items = content["items"]
            eng = items["engineer"]
            self.assertEqual(eng["include"], {"full": True, "professional": True, "resume": False})
            self.assertEqual(eng["role"], {"en": "Engineer", "es": "Ingeniero"})
            self.assertEqual(eng["org"], {"en": "Org, Town", "es": "Org, Pueblo"})
            self.assertEqual(eng["order"]["full"], ["engineer-1", "engineer-2"])
            self.assertEqual(eng["order"]["professional"], ["engineer-1"])
            self.assertEqual(eng["children"]["engineer-2"]["include"], {"full": True, "professional": False, "resume": False})
            self.assertEqual(items["tutor"]["include"]["professional"], False)
            exp = next(s for s in content["sections"] if s["id"] == "experience")
            self.assertEqual(exp["order"], {"full": ["engineer", "tutor"], "professional": ["engineer"]})
            self.assertEqual(content["header"]["title"]["full"], {"en": "Engineer · Tester", "es": "Ingeniero · Probador"})
            self.assertEqual(report["variants"]["professional"]["sharedWithEarlier"], 4)  # education entry, program line, engineer entry, its first bullet
            self.assertEqual(report["unpaired"], {"full": [], "professional": []})
            self.assertEqual(len(slots["Gideon Ong CV Full EN.docx"]), 12)  # one slot per paragraph
            self.assertIn("engineer", content["hashes"]["Gideon Ong CV Full EN.docx"])

    def test_pairing_gap_and_proof(self):
        with TempRepo() as t:
            en = make_master(t.root / "Gideon Ong Resume EN.docx", "en", title=False, entries=[("Engineer", "Org", "Summer 2026", ["One.", "Two."])])
            es = make_master(t.root / "Gideon Ong Resume ES.docx", "es", title=False, entries=[("Ingeniero", "Org", "Verano 2026", ["Uno."])], split_runs=False)
            docs = {"resume": {"en": importer.read_master(en), "es": importer.read_master(es)}}
            content, slots, report = importer.merge(docs)
            self.assertEqual(len(report["unpaired"]["resume"]), 1)
            self.assertIn("no ES twin", report["unpaired"]["resume"][0])
            self.assertEqual(content["items"]["engineer"]["children"]["engineer-2"]["text"], {"en": "Two.", "es": ""})
            self.assertIsNone(content["header"]["title"]["resume"])
            tmp = Path(tempfile.mkdtemp())
            for doc in (docs["resume"]["en"], docs["resume"]["es"]):
                out = renderer.force_render(doc, slots[doc.path.name], content, tmp / doc.path.name)
                self.assertEqual(importer.compare(doc.path, out), [], doc.path.name)
            # an edited render changes exactly that paragraph and re-imports to the new text
            c2 = copy.deepcopy(content)
            c2["items"]["engineer"]["children"]["engineer-1"]["text"]["en"] = "One, edited."
            c2["items"]["engineer"]["org"]["en"] = "New Org"
            doc = importer.read_master(en)
            out = renderer.force_render(doc, slots[en.name], c2, tmp / "edited.docx")
            problems = importer.compare(en, out)
            self.assertEqual(len(problems), 2, problems)
            nodes = importer.structure(importer.read_master(out))
            self.assertIn({"text": "One, edited."}, [n.parts for n in nodes])
            e = next(n.parts for n in nodes if n.kind == "entry" and n.parts["role"] == "Engineer")
            self.assertEqual((e["org"], e["date"], e["sep"]), ("New Org", "Summer 2026", " · "))
            self.assertEqual(importer.compare(en, en), [])

    def test_service_import_writes_only_when_lossless(self):
        with TempRepo() as t:
            (t.root / masters.MASTERS_DIR).mkdir(parents=True)
            for m in masters.MASTERS:
                make_master(t.root / masters.MASTERS_DIR / m.file, m.lang, title=m.variant != "resume", split_runs=(m.lang == "en"))
            svc = CvTextService(t.root, Backups(t.root, t.local / "backups"))
            self.assertIsNone(svc.state()["content"])
            r = svc.run_import()
            self.assertTrue(r["ok"], r)
            self.assertTrue(r["written"])
            self.assertEqual({k: v for k, v in r["proof"].items()}, {m.file: [] for m in masters.MASTERS})
            self.assertTrue(svc.content_path().is_file())
            self.assertEqual(len(list((t.root / "staging/cv-content/slots").glob("*.json"))), 6)
            st = svc.state()
            self.assertEqual(st["content"]["sections"], 2)
            self.assertTrue(st["lastImport"]["ok"])
            # re-import backs the previous set up
            r2 = svc.run_import()
            self.assertTrue(r2["ok"])
            self.assertIn("backup", r2)
            # a master with an unsupported construct blocks the import
            from docx import Document

            p = t.root / masters.MASTERS_DIR / "Gideon Ong Resume ES.docx"
            doc = Document(str(p))
            doc.add_table(rows=1, cols=1)
            doc.save(str(p))
            r3 = svc.run_import()
            self.assertEqual((r3["ok"], r3["error"]), (False, "unsupported"))
            p.unlink()
            self.assertEqual(svc.run_import()["error"], "missing-masters")


@unittest.skipUnless(MASTERS_PRESENT, "CV masters are not on this machine")
class RealMastersTests(unittest.TestCase):
    def test_all_six_round_trip_losslessly(self):
        docs = {v: {lang: importer.read_master(masters.docx_path(REPO_ROOT, masters.BY_ID[f"{v}-{lang}"])) for lang in ("en", "es")} for v in importer.VARIANTS}
        content, slots, report = importer.merge(docs)
        self.assertGreater(len(content["items"]), 30)
        self.assertEqual(report["unpaired"]["full"], [])
        self.assertEqual(report["unpaired"]["professional"], [])
        self.assertEqual(len(report["unpaired"]["resume"]), 1)  # the ES résumé lacks the IEL line
        tmp = Path(tempfile.mkdtemp())
        for langs in docs.values():
            for doc in langs.values():
                out = renderer.force_render(doc, slots[doc.path.name], content, tmp / doc.path.name)
                self.assertEqual(importer.compare(doc.path, out), [], doc.path.name)
        self.assertEqual(report["formatting"], [])


if __name__ == "__main__":
    unittest.main()
