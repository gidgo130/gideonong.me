"""Phase 1b: CV / résumé check & publish.

Synthetic .docx / .pdf files carry made-up text, so nothing private is in the
test code. Tests that look at the real masters (gitignored) or the committed
public PDFs skip when those files are absent. The Word export runs only with
EDITOR_WORD_TESTS=1 (it starts a private Word instance).
"""

import os
import time
import unittest
from pathlib import Path

from _helpers import REPO_ROOT, TempRepo
from core.backups import Backups
from core.cv import docxread, export, fit, masters, pdfcheck, publish, scans
from core.cv.service import CvService

PROFESSIONAL_EN = REPO_ROOT / masters.MASTERS_DIR / "Gideon Ong CV Professional EN.docx"
MASTERS_PRESENT = PROFESSIONAL_EN.is_file()
PUBLIC_CV_PDFS = sorted((REPO_ROOT / "assets" / "pdfs" / "cv").glob("en Gideon Ong CV *.pdf")) if (REPO_ROOT / "assets" / "pdfs" / "cv").is_dir() else []
PUBLIC_CV_PDF = PUBLIC_CV_PDFS[-1] if PUBLIC_CV_PDFS else None


# ----------------------------------------------------------------- builders


def make_docx(path: Path, entries, bullets=(), margin_in=0.5, tab_in=7.5, extra=()):
    """A tiny CV-shaped document: name, heading, entry lines with a right tab, bullets."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(margin_in)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Georgia", Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Test Person")
    r.bold, r.font.size = True, Pt(18)
    p = doc.add_paragraph()
    r = p.add_run("Experience")
    r.bold, r.font.size = True, Pt(13)
    for role, org, date in entries:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(tab_in), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(role)
        r.bold = True
        p.add_run(" · " + org)
        p.add_run("\t")
        r = p.add_run(date.replace(" ", "\xa0"))
        r.italic = r.underline = True
        for b in bullets:
            doc.add_paragraph(b, style="List Bullet")
    for t in extra:
        doc.add_paragraph(t)
    doc.save(str(path))
    return path


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages):
    """A valid single-font PDF; `pages` is a list of pages, each a list of text lines (Latin-1)."""
    objects = []
    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, lines in enumerate(pages):
        parts = ["BT /F1 11 Tf 72 720 Td 14 TL"]
        for line in lines:
            parts.append(f"({_pdf_escape(line)}) Tj T*")
        parts.append("ET")
        content = " ".join(parts).encode("cp1252")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for o in offsets:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(out))
    return path


_TICK = [0]


def _later(path: Path, than: Path) -> None:
    """Give `path` an mtime after `than`, later on every call (file systems round; tests must not depend on luck)."""
    _TICK[0] += 5
    t = os.stat(than).st_mtime + _TICK[0]
    os.utime(path, (t, t))


ENTRY = ("Test Engineer", "Example Org, Somewhere", "Summer 2026")
BULLETS = ("Built a thing that saved 10,000 hours.", "Wrote the guide in two days.")
DOC_LINES = ["Test Person", "Experience", "Test Engineer · Example Org, Somewhere Summer 2026", *BULLETS]


class FakeRepo:
    """TempRepo plus synthetic masters + PDFs for the four publishable documents."""

    def __init__(self, t: TempRepo):
        self.t = t
        self.root = t.root
        self.mdir = t.root / masters.MASTERS_DIR
        self.mdir.mkdir(parents=True)
        self.out = t.root / masters.OUT_DIR
        self.out.mkdir(parents=True)
        for m in masters.MASTERS:
            if not m.publish_type:
                continue
            make_docx(self.mdir / m.file, [ENTRY], BULLETS)
        for m in masters.MASTERS:
            if not m.publish_type:
                continue
            self.write_pdf(m, [DOC_LINES])
        self.backups = Backups(t.root, t.local / "backups")
        self.service = CvService(t.root, self.backups)

    def write_pdf(self, m, pages):
        p = make_pdf(masters.pdf_path(self.root, m), pages)
        _later(p, masters.docx_path(self.root, m))
        return p


# ----------------------------------------------------------------- tests


class MastersTests(unittest.TestCase):
    def test_table_matches_manifest_naming(self):
        for m in masters.MASTERS:
            if m.publish_type:
                self.assertRegex(publish.target_rel(m, "20260928"), r"^assets/pdfs/(cv|resume)/(en|es) Gideon Ong (CV|Resume) 20260928\.pdf$")
                self.assertTrue(publish.published_pattern(m).match(Path(publish.target_rel(m, "20260928")).name))
        self.assertEqual(masters.BY_ID["resume-en"].pair_id, "resume-es")
        self.assertTrue(masters.BY_ID["resume-es"].one_page)
        self.assertFalse(masters.BY_ID["professional-en"].one_page)

    def test_word_lock_detection(self):
        with TempRepo() as t:
            d = t.root / "staging" / "x"
            d.mkdir(parents=True)
            docx = d / "Gideon Ong Resume EN.docx"
            docx.write_bytes(b"")
            self.assertIsNone(masters.word_lock_file(docx))
            self.assertIsNone(export.check_exportable(docx))
            (d / "~$deon Ong Resume EN.docx").write_bytes(b"")
            self.assertIsNotNone(masters.word_lock_file(docx))
            self.assertIn("close it in Word", export.check_exportable(docx))
            self.assertIn("missing", export.check_exportable(d / "nope.docx"))
            with self.assertRaises(export.ExportError):
                export.export_pdf(docx, d / "out.pdf")

    def test_info_reports_stale_pdf(self):
        with TempRepo() as t:
            f = FakeRepo(t)
            m = masters.BY_ID["resume-en"]
            info = masters.info(t.root, m)
            self.assertTrue(info["exists"] and info["pdf"]["exists"] and not info["pdf"]["stale"])
            _later(masters.docx_path(t.root, m), masters.pdf_path(t.root, m))
            self.assertTrue(masters.info(t.root, m)["pdf"]["stale"])


class DocxReadAndFitTests(unittest.TestCase):
    def test_kinds_fonts_and_fit_on_a_synthetic_doc(self):
        with TempRepo() as t:
            path = make_docx(t.root / "a.docx", [ENTRY], BULLETS)
            d = docxread.read(path)
            self.assertEqual([p.kind for p in d.paragraphs], ["name", "heading", "entry", "bullet", "bullet"])
            self.assertAlmostEqual(d.text_width_pt, 7.5 * 72, places=1)
            self.assertEqual(d.body_size_pt, 11.0)
            entry = d.entries()[0]
            self.assertEqual([(r.bold, r.italic) for r in entry.runs], [(True, False), (False, False), (False, False), (False, True)])
            self.assertAlmostEqual(entry.tab_right_pt, 540.0, places=1)
            self.assertGreater(d.paragraphs[3].left_indent_pt, 0)  # bullets are indented by their list level
            line = fit.check_entry(entry, d)
            self.assertEqual((line.status, line.text, line.date), ("fits", "Test Engineer · Example Org, Somewhere", "Summer 2026"))
            self.assertAlmostEqual(line.avail_pt, 540.0, places=1)
            self.assertAlmostEqual(line.used_pt, fit.width_pt("Test Engineer", bold=True) + fit.width_pt(" · Example Org, Somewhere") + fit.width_pt(" ") + fit.width_pt("Summer 2026", italic=True), places=3)

    def test_overflow_and_tight(self):
        with TempRepo() as t:
            date = "Summer 2026"
            fixed = fit.width_pt(" · Org") + fit.width_pt(" ") + fit.width_pt(date, italic=True)
            long_role = "Very long role title " * 6
            role = "a"
            while fit.width_pt(role, bold=True) + fixed < 540 * 0.985:
                role += "a"
            path = make_docx(t.root / "b.docx", [(long_role, "Org", date), (role, "Org", date), ("Short", "Org", date)])
            lines = fit.check_doc(docxread.read(path))
            self.assertEqual([l.status for l in lines], ["overflow", "tight", "fits"])
            self.assertLess(lines[0].left_pt, 0)
            s = fit.summary(lines)
            self.assertEqual((s["count"], s["overflow"], s["tight"]), (3, 1, 1))
            self.assertEqual(s["worst"]["index"], lines[0].index)

    def test_measurement_is_subpixel(self):
        # FreeType rounds advances to whole pixels at the nominal size; the 64× trick must not
        self.assertNotAlmostEqual(fit.width_pt("Summer 2026", italic=True) % 1, 0.0, places=2)
        self.assertEqual(fit.width_pt(""), 0.0)
        self.assertEqual(fit.width_pt("a\xa0b"), fit.width_pt("a b"))

    @unittest.skipUnless(PUBLIC_CV_PDF, "no public CV PDF in assets/pdfs/cv")
    def test_measurement_agrees_with_word(self):
        pdf_w = pdfcheck.word_width(PUBLIC_CV_PDF, "Summer 2026")
        self.assertIsNotNone(pdf_w)
        self.assertAlmostEqual(pdf_w, fit.width_pt("Summer 2026", italic=True, size_pt=11), delta=1.0)

    @unittest.skipUnless(MASTERS_PRESENT, "CV masters are not on this machine")
    def test_real_masters_read_and_pair(self):
        docs = {m.id: docxread.read(masters.docx_path(REPO_ROOT, m)) for m in masters.MASTERS if masters.docx_path(REPO_ROOT, m).is_file()}
        self.assertEqual(len(docs["professional-en"].entries()), 21)
        for variant in ("full", "professional", "resume"):
            en, es = docs.get(f"{variant}-en"), docs.get(f"{variant}-es")
            if en and es:
                self.assertEqual(len(en.entries()), len(es.entries()), variant)
        for d in docs.values():
            self.assertEqual([l.status for l in fit.check_doc(d) if l.status == "overflow"], [])


class PdfCheckTests(unittest.TestCase):
    def test_lines_pages_and_hyphen_tolerance(self):
        with TempRepo() as t:
            pdf = make_pdf(t.root / "x.pdf", [["Alpha Role · Beta Org Fall 2025", "We wrote lessons-", "learned documents."], ["Second page"]])
            r = pdfcheck.read_pdf(pdf)
            self.assertEqual(r.pages, 2)
            self.assertEqual(r.lines[0], "Alpha Role · Beta Org Fall 2025")
            self.assertEqual(r.page_of_line[-1], 2)
            self.assertEqual(pdfcheck.entry_lines_check(["Alpha Role · Beta Org\tFall\xa02025"], r), [{"text": "Alpha Role · Beta Org Fall 2025", "ok": True}])
            self.assertEqual(pdfcheck.entry_lines_check(["Alpha Role · Beta Org\tSpring 2025"], r)[0]["ok"], False)
            self.assertEqual(pdfcheck.coverage_check(["We wrote lessons-learned documents.", "Second page", ""], r), [])
            self.assertEqual(pdfcheck.coverage_check(["Not in there"], r), ["Not in there"])

    @unittest.skipUnless(MASTERS_PRESENT and PUBLIC_CV_PDF, "needs the Professional EN master and the public CV PDF")
    def test_published_cv_matches_its_master(self):
        d = docxread.read(PROFESSIONAL_EN)
        r = pdfcheck.read_pdf(PUBLIC_CV_PDF)
        self.assertGreaterEqual(r.pages, 1)
        self.assertTrue(all(e["ok"] for e in pdfcheck.entry_lines_check([p.text for p in d.entries()], r)))
        self.assertEqual(pdfcheck.coverage_check([p.text for p in d.nonempty()], r), [])


class ScanTests(unittest.TestCase):
    def test_blocklist(self):
        with TempRepo() as t:
            p = t.root / "staging" / "editor-private" / "blocklist.txt"
            terms, warn = scans.load_blocklist(p)
            self.assertEqual(terms, [])
            self.assertIn("no blocklist", warn)
            p.parent.mkdir(parents=True)
            p.write_text("# comment\nSecret  Widget | internal name\n\n555-0100\n", encoding="utf-8")
            terms, warn = scans.load_blocklist(p)
            self.assertIsNone(warn)
            self.assertEqual([(x.text, x.reason) for x in terms], [("Secret Widget", "internal name"), ("555-0100", "")])
            hits = scans.blocklist_hits("Built the secret\twidget for fun", terms)
            self.assertEqual([h["term"] for h in hits], ["Secret Widget"])
            self.assertIn("secret widget", hits[0]["context"].lower())
            self.assertEqual(scans.blocklist_hits("nothing here", terms), [])
            # decision 9: the (918) number is stripped before matching on published PDFs
            phone = [scans.Term("409-2444", "phone", 1)]
            self.assertEqual(len(scans.blocklist_hits("call +1 (918) 409-2444", phone)), 1)
            self.assertEqual(scans.blocklist_hits("call +1 (918) 409-2444", phone, scans.ALLOWED_IN_PUBLISHED_PDFS), [])

    def test_voice_words_with_spinelli_exemption(self):
        with TempRepo() as t:
            path = make_docx(t.root / "v.docx", [("Translator", "Spinelli Institute", "2021"), ("Other", "Org", "2020")], ["Leveraged a passionate approach."])
            hits = scans.voice_hits(docxread.read(path))
            self.assertEqual([h["words"] for h in hits], [["passionate"], ["leveraged", "passionate"]])

    def test_parity(self):
        with TempRepo() as t:
            en = docxread.read(make_docx(t.root / "en.docx", [("Role", "Org", "Summer 2026")], ["Saved $10,000 a year.", "Second bullet."]))
            es = docxread.read(make_docx(t.root / "es.docx", [("Puesto", "Org", "Verano de 2026")], ["Ahorré US$10k al año.", "Segundo punto."]))
            self.assertEqual(scans.parity(en, es), [])
            es2 = docxread.read(make_docx(t.root / "es2.docx", [("Puesto", "Org", "Verano de 2025")], ["Ahorré US$12 mil al año."]))
            notes = scans.parity(en, es2)
            self.assertEqual(len(notes), 3)  # paragraph count, entry numbers, bullet count
            self.assertTrue(any("numbers differ" in n for n in notes))
            self.assertTrue(any("bullets" in n for n in notes))
            es3 = docxread.read(make_docx(t.root / "es3.docx", [("A", "B", "2026"), ("C", "D", "2025")], ["x"]))
            self.assertTrue(any("entry lines" in n for n in scans.parity(en, es3)))


class ServiceAndPublishTests(unittest.TestCase):
    def test_checks_and_publish_flow(self):
        with TempRepo() as t:
            f = FakeRepo(t)
            old_en = t.root / "assets/pdfs/resume/en Gideon Ong Resume 20250101.pdf"
            old_en.write_bytes(b"%PDF-old")
            (t.root / "assets/pdfs/cv/en Gideon Ong CV 20260928.pdf").write_bytes(b"%PDF-cv-same-date")
            checks = f.service.check()
            self.assertEqual(set(checks), {m.id for m in masters.MASTERS})
            self.assertEqual(checks["full-en"]["errors"][0]["code"], "no-master")
            for mid in ("resume-en", "resume-es", "professional-en", "professional-es"):
                self.assertEqual(checks[mid]["errors"], [], mid)
                self.assertEqual(checks[mid]["pdf"]["pages"], 1)
                self.assertTrue(all(e["ok"] for e in checks[mid]["pdf"]["entryLines"]))
                self.assertEqual(checks[mid]["fit"]["summary"]["count"], 1)
            self.assertTrue(any(w["code"] == "no-blocklist" for w in checks["resume-en"]["warnings"]))
            # plan: the résumé adds and moves the old dated file; the CV overwrites a same-date file
            p = f.service.publish_plan("20260928")
            self.assertTrue(p["ok"], p["errors"])
            by_id = {i["id"]: i for i in p["items"]}
            self.assertEqual(by_id["resume-en"]["action"], "add")
            self.assertEqual(by_id["resume-en"]["moves"], ["assets/pdfs/resume/en Gideon Ong Resume 20250101.pdf"])
            self.assertEqual(by_id["professional-en"]["action"], "overwrite")
            self.assertEqual(by_id["professional-es"]["moves"], [])
            self.assertFalse(f.service.publish_plan("2026-09-28")["ok"])
            # apply
            r = f.service.publish_apply("20260928")
            self.assertTrue(r["ok"], r)
            self.assertEqual(sorted(r["added"]), sorted(["assets/pdfs/resume/en Gideon Ong Resume 20260928.pdf", "assets/pdfs/resume/es Gideon Ong Resume 20260928.pdf", "assets/pdfs/cv/es Gideon Ong CV 20260928.pdf"]))
            self.assertEqual(r["overwritten"], ["assets/pdfs/cv/en Gideon Ong CV 20260928.pdf"])
            self.assertEqual(r["moved"], ["assets/pdfs/resume/en Gideon Ong Resume 20250101.pdf"])
            self.assertFalse(old_en.exists())
            self.assertEqual((t.root / r["added"][0]).read_bytes()[:5], b"%PDF-")
            self.assertNotEqual((t.root / "assets/pdfs/cv/en Gideon Ong CV 20260928.pdf").read_bytes(), b"%PDF-cv-same-date")
            bset = f.backups.get(r["backup"])
            # (js/docs-data.js did not exist in the temp repo before the publish, so it is not in the set)
            self.assertEqual({x["path"] for x in bset.files}, {"assets/pdfs/resume/en Gideon Ong Resume 20250101.pdf", "assets/pdfs/cv/en Gideon Ong CV 20260928.pdf"})
            self.assertEqual(f.backups.read_file(r["backup"], "assets/pdfs/resume/en Gideon Ong Resume 20250101.pdf"), b"%PDF-old")
            if r["manifest"]["ran"] and r["manifest"]["ok"]:
                text = (t.root / "js" / "docs-data.js").read_text(encoding="utf-8")
                self.assertIn('resume: { en: "assets/pdfs/resume/en Gideon Ong Resume 20260928.pdf"', text)
                self.assertIn("transcript: { en: null, es: null }", text)
            else:
                self.assertIn("node", r["manifest"]["output"])
            # a second publish on the same date overwrites everything and moves nothing
            p2 = f.service.publish_plan("20260928")
            self.assertTrue(all(i["action"] == "overwrite" and not i["moves"] for i in p2["items"]))

    def test_gates_block_publish(self):
        with TempRepo() as t:
            f = FakeRepo(t)
            m = masters.BY_ID["resume-en"]
            # two pages
            f.write_pdf(m, [DOC_LINES, ["overflow"]])
            c = f.service.check()["resume-en"]
            self.assertIn("pages", [e["code"] for e in c["errors"]])
            # the CV may have two pages
            f.write_pdf(masters.BY_ID["professional-en"], [DOC_LINES, ["more"]])
            self.assertNotIn("pages", [e["code"] for e in f.service.check()["professional-en"]["errors"]])
            # a wrapped entry line
            f.write_pdf(m, [["Test Person", "Experience", "Test Engineer · Example Org, Somewhere", "Summer 2026", *BULLETS]])
            self.assertIn("wrapped", [e["code"] for e in f.service.check()["resume-en"]["errors"]])
            # a paragraph edited after the export (text missing) and an older PDF
            f.write_pdf(m, [[*DOC_LINES[:3], "Built a thing that saved 9,000 hours.", DOC_LINES[4]]])
            self.assertIn("stale-text", [e["code"] for e in f.service.check()["resume-en"]["errors"]])
            f.write_pdf(m, [DOC_LINES])
            self.assertEqual(f.service.check()["resume-en"]["errors"], [])
            _later(masters.docx_path(t.root, m), masters.pdf_path(t.root, m))
            self.assertEqual([e["code"] for e in f.service.check()["resume-en"]["errors"]], ["stale"])
            f.write_pdf(m, [DOC_LINES])
            self.assertEqual(f.service.check()["resume-en"]["errors"], [])
            # blocklist hit
            bl = t.root / masters.BLOCKLIST_PATH
            bl.parent.mkdir(parents=True, exist_ok=True)
            bl.write_text("Example Org | client name\n", encoding="utf-8")
            c = f.service.check()["resume-en"]
            self.assertEqual([e["code"] for e in c["errors"]], ["blocklist"])
            self.assertIn("client name", c["errors"][0]["message"])
            r = f.service.publish_apply("20260928")
            self.assertEqual((r["ok"], r["error"]), (False, "blocked"))
            self.assertEqual(list((t.root / "assets/pdfs/resume").iterdir()), [])
            self.assertEqual(f.backups.list(), [])
            # cache: the same inputs are not re-read; a changed blocklist is
            bl.write_text("", encoding="utf-8")
            _later(bl, masters.pdf_path(t.root, m))
            self.assertEqual(f.service.check()["resume-en"]["errors"], [])

    def test_export_job_refuses_open_masters_and_runs_in_background(self):
        with TempRepo() as t:
            f = FakeRepo(t)
            m = masters.BY_ID["resume-en"]
            (f.mdir / ("~$" + m.file[2:])).write_bytes(b"")
            with self.assertRaises(RuntimeError):
                f.service.start_export(["resume-en"])
            calls = []

            def fake_export(docx, pdf, log):
                calls.append(Path(docx).name)
                log("fake export")
                Path(pdf).write_bytes(b"%PDF-fake")

            import core.cv.service as service_mod

            orig = service_mod.export.export_pdf
            service_mod.export.export_pdf = fake_export
            try:
                job = f.service.start_export(["resume-en", "resume-es"])
                self.assertEqual(job["ids"], ["resume-es"])
                self.assertIn("resume-en", job["refused"])
                for _ in range(100):
                    if not f.service.busy():
                        break
                    time.sleep(0.05)
            finally:
                service_mod.export.export_pdf = orig
            j = f.service.job_json()
            self.assertFalse(j["running"])
            self.assertEqual((j["done"], j["failed"]), (["resume-es"], {}))
            self.assertTrue(any("fake export" in line for line in j["log"]))
            self.assertEqual(calls, [masters.BY_ID["resume-es"].file])

    @unittest.skipUnless(os.environ.get("EDITOR_WORD_TESTS") == "1", "set EDITOR_WORD_TESTS=1 to run the Word export")
    def test_word_export_for_real(self):
        with TempRepo() as t:
            path = make_docx(t.root / "w.docx", [ENTRY], BULLETS)
            pdf = export.export_pdf(path, t.root / "w.pdf")
            r = pdfcheck.read_pdf(pdf)
            self.assertEqual(r.pages, 1)
            self.assertTrue(pdfcheck.entry_lines_check([docxread.read(path).entries()[0].text], r)[0]["ok"])


if __name__ == "__main__":
    unittest.main()
