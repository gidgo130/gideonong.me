"""Build Gideon's public "Unofficial Academic Record" (English + Spanish) from
transcript-data.json + course-titles.json + profile.json.

Usage:
    python make_transcript.py              # .docx files into scripts/transcript/output/ (for checking)
    python make_transcript.py --pdf        # also PDFs via Word into assets/pdfs/transcript/
    python make_transcript.py --pdf --date 20260914 --pdf-dir SOME/DIR
Needs: python-docx (pip install python-docx); pywin32 only with --pdf (a private, hidden Word
instance does the export — your own open Word is never touched or closed).

Safety checks (the build stops if any fails):
  * every course code has a title in course-titles.json;
  * recomputed term and cumulative credits / grade points / GPA equal TU's printed totals.
Titles marked "verified": false are listed as a warning so they can be checked against
bulletin.utulsa.edu; use --strict to make that an error.
"""
import argparse, json, os, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = 'Georgia'
POINTS = {'A': 4.0, 'B': 3.0, 'C': 2.0, 'D': 1.0, 'F': 0.0}
NO_GPA = {'P', 'S', 'U', 'W', 'I', 'NG', ''}

T = {  # interface text, EN / ES
  'en': dict(title='Unofficial Academic Record', program='Program', majors='Majors', minors='Minors',
             grad='Expected graduation', gpa='Cumulative GPA', hours='Total credit hours',
             asof='Record current as of {date}',
             exam='Credit by Examination', exam_kinds={'Proficiency': 'Advanced Placement (AP) Equivalencies', 'AP': 'Advanced Placement (AP) Equivalencies'},
             cols=('Course', 'Title', 'Credits', 'Grade'), in_progress='In progress',
             term_line='Term: {cr} credits · GPA {gpa} · Cumulative GPA {cum}',
             term_line_nogpa='Term: {cr} credits (pass/fail only)',
             term_line_ip='{cr} credits in progress',
             seasons={'Fall': 'Fall', 'Spring': 'Spring', 'Summer': 'Summer', 'Winter': 'Winter'},
             key='Grades: A = 4.0 (highest), B = 3.0, C = 2.0, D = 1.0, F = 0. P = pass: credit earned, not counted in the GPA. GPA covers work taken at the University of Tulsa only.',
             footer='Unofficial record prepared by {name} from his University of Tulsa unofficial transcript printed {date}. An official transcript can be requested through the University of Tulsa Office of the Registrar.',
             months=['January','February','March','April','May','June','July','August','September','October','November','December'],
             date_fmt='{month} {day}, {year}'),
  'es': dict(title='Expediente académico no oficial', program='Programa', majors='Carreras', minors='Especializaciones menores',
             grad='Graduación prevista', gpa='Promedio general (GPA)', hours='Total de créditos',
             asof='Expediente actualizado al {date}',
             exam='Créditos por examen', exam_kinds={'Proficiency': 'Equivalencias de Advanced Placement (AP)', 'AP': 'Equivalencias de Advanced Placement (AP)'},
             cols=('Curso', 'Título', 'Créditos', 'Calificación'), in_progress='En curso',
             term_line='Periodo: {cr} créditos · promedio {gpa} · promedio general {cum}',
             term_line_nogpa='Periodo: {cr} créditos (solo aprobado/no aprobado)',
             term_line_ip='{cr} créditos en curso',
             seasons={'Fall': 'Otoño', 'Spring': 'Primavera', 'Summer': 'Verano', 'Winter': 'Invierno'},
             key='Escala de calificaciones estadounidense: A = 4.0 (la más alta), B = 3.0, C = 2.0, D = 1.0, F = 0. P = aprobado: otorga créditos, pero no cuenta en el promedio. El promedio incluye solo los cursos tomados en la Universidad de Tulsa.',
             footer='Expediente no oficial preparado por {name} a partir de su historial académico no oficial de la Universidad de Tulsa, impreso el {date}. El historial académico oficial se puede solicitar a la Oficina del Registrador (Office of the Registrar) de la Universidad de Tulsa.',
             months=['enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre'],
             date_fmt='{day} de {month} de {year}'),
}

def load(name):
    with open(os.path.join(HERE, name), encoding='utf-8') as f:
        return json.load(f)

def fmt_date(iso, lang):
    y, m, d = (int(x) for x in iso.split('-'))
    t = T[lang]
    return t['date_fmt'].format(month=t['months'][m - 1], day=d, year=y)

def title_for(course, titles):
    entry = titles.get(course['code'])
    if entry is None:
        return None
    if 'sections' in entry:
        return entry['sections'].get(course['transcript_title'])
    return entry

# ---------------------------------------------------------------- checks
def apply_adjustments(data, adj):
    """Insert manual additions (courses not printed on the TU transcript). Returns warning lines."""
    notes = []
    for a in adj.get('add', []):
        block = next(b for b in data['blocks'] if b.get('source') == a['block'] or f"{b.get('season')} {b.get('year')}" == a['block'])
        codes = [c['code'] for c in block['courses']]
        if a['course']['code'] in codes:
            notes.append(f"{a['course']['code']} is now on the TU transcript - remove it from adjustments.json")
            continue
        pos = codes.index(a['after']) + 1 if a.get('after') in codes else len(codes)
        block['courses'].insert(pos, dict(a['course']))
        notes.append(f"added {a['course']['code']} to {a['block']} (not on TU transcript): {a.get('reason', '')}")
    return notes

def totals(data):
    cum = dict(earned=0, gpa_hours=0, points=0.0)
    for b in data['blocks']:
        done = [c for c in b['courses'] if c['grade']]
        term = dict(earned=sum(c['credits'] for c in done),
                    gpa_hours=sum(c['credits'] for c in done if c['grade'] in POINTS),
                    points=sum(POINTS[c['grade']] * c['credits'] for c in done if c['grade'] in POINTS))
        for k in cum:
            cum[k] += term[k]
        b['_term'], b['_cum'] = term, dict(cum)
    return cum

def check(data, titles, strict):
    problems, unverified = [], []
    for b in data['blocks']:
        for c in b['courses']:
            t = title_for(c, titles)
            if t is None:
                problems.append(f"no title for {c['code']} ({c['transcript_title']}) in course-titles.json")
            elif not t.get('verified', False):
                unverified.append(f"{c['code']} {t['en']}  [{t.get('source', '?')}]")
            if c['grade'] and c['grade'] not in POINTS and c['grade'] not in NO_GPA:
                problems.append(f"unknown grade {c['grade']!r} for {c['code']}")
    cum = dict(earned=0, gpa_hours=0, points=0.0)
    for b in data['blocks']:
        done = [c for c in b['courses'] if c['grade']]
        term = dict(earned=sum(c['credits'] for c in done),
                    gpa_hours=sum(c['credits'] for c in done if c['grade'] in POINTS),
                    points=sum(POINTS[c['grade']] * c['credits'] for c in done if c['grade'] in POINTS))
        for k in cum:
            cum[k] += term[k]
        b['_term'], b['_cum'] = term, dict(cum)
        label = b.get('source') or f"{b['season']} {b['year']}"
        for mine, key in ((term, 'term_totals'), (cum, 'cumulative_totals')):
            tu = b.get(key)
            if not tu or b['kind'] == 'exam' and key == 'term_totals':
                continue  # TU prints the exam totals once, for both exam groups together
            if (tu['earned'], tu['gpa_hours']) != (mine['earned'], mine['gpa_hours']) or abs(tu['grade_points'] - mine['points']) > 0.01:
                problems.append(f"{label} {key}: TU {tu} vs computed {mine}")
    exam_total = sum(c['credits'] for b in data['blocks'] if b['kind'] == 'exam' for c in b['courses'])
    last_exam = [b for b in data['blocks'] if b['kind'] == 'exam'][-1]
    if last_exam.get('term_totals') and last_exam['term_totals']['earned'] != exam_total:
        problems.append(f"exam credit: TU {last_exam['term_totals']['earned']} vs computed {exam_total}")
    for u in unverified:
        print('unverified title:', u)
    if strict and unverified:
        problems.append(f'{len(unverified)} unverified titles (--strict)')
    if problems:
        sys.exit('STOPPED:\n  ' + '\n  '.join(problems))
    return cum, exam_total

# ---------------------------------------------------------------- docx helpers
def style_run(r, size=11, bold=False, italic=False, color=None):
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r

def para(doc, text='', size=11, bold=False, italic=False, align=None, before=0, after=0, color=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.15
    pf.keep_with_next = keep
    if align:
        p.alignment = align
    if text:
        style_run(p.add_run(text), size, bold, italic, color)
    return p

def label_line(doc, label, value):
    p = para(doc)
    style_run(p.add_run(label + ': '), bold=True)
    style_run(p.add_run(value))

def insert_before(parent, el, names):
    """Insert el before the first child named in names (keeps the OOXML schema order)."""
    for child in parent:
        if child.tag in {qn(n) for n in names}:
            child.addprevious(el); return
    parent.append(el)

def rule(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', '000000')):
        bottom.set(qn(k), v)
    b.append(bottom)
    insert_before(pPr, b, ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:spacing', 'w:ind', 'w:jc', 'w:rPr'))

def heading(doc, text):
    p = para(doc, text, size=13, bold=True, before=10, after=3, keep=True)
    rule(p)

def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'nil'); borders.append(e)
    insert_before(tblPr, borders, ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption'))

def cell_text(cell, text, bold=False, italic=False, align=None, size=10.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
    if align:
        p.alignment = align
    style_run(p.add_run(text), size=size, bold=bold, italic=italic)

WIDTHS = (Inches(1.05), Inches(4.4), Inches(0.9), Inches(1.15))

def course_table(doc, courses, titles, lang, in_progress=False):
    t = T[lang]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    no_borders(table)
    for i, h in enumerate(t['cols']):
        cell_text(table.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER if i >= 2 else None)
    for c in courses:
        row = table.add_row().cells
        grade = t['in_progress'] if in_progress else (c['grade'] or '—')
        cell_text(row[0], c['code'])
        cell_text(row[1], title_for(c, titles)[lang])
        cell_text(row[2], str(c['credits']), align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[3], grade, italic=in_progress, align=WD_ALIGN_PARAGRAPH.CENTER)
    # fixed layout + explicit grid widths so Word, LibreOffice and iPad agree on the columns
    tblPr = table._tbl.tblPr
    # (table.autofit = False above already sets <w:tblLayout w:type="fixed"/>)
    for i, w in enumerate(WIDTHS):
        table.columns[i].width = w
    for row in table.rows:
        for i, w in enumerate(WIDTHS):
            row.cells[i].width = w
        tr = row._tr.get_or_add_trPr(); cant = OxmlElement('w:cantSplit'); tr.append(cant)
    # repeat header row if a table crosses pages; keep rows together with the heading
    hdr = table.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); hdr.append(h)
    for row in table.rows[:2]:  # keep the header row with the first course (tables may split after that)
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True
    return table

def gpa(points, hours):
    return f'{points / hours:.2f}' if hours else '—'

def export_pdf(docx_path, pdf_path):
    """docx → PDF through a PRIVATE Word instance (DispatchEx), opened read-only, quit afterwards.
    Not docx2pdf.convert(): that attaches to the Word already running and calls Quit() on it, which
    would close whatever else is open in Word (same fix as scripts/editor/core/cv/export.py)."""
    import pythoncom, win32com.client  # pywin32
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx('Word.Application')
        word.Visible, word.DisplayAlerts = False, 0
        doc = word.Documents.Open(os.path.abspath(docx_path), ReadOnly=True, AddToRecentFiles=False, Visible=False)
        try:
            doc.ExportAsFixedFormat(os.path.abspath(pdf_path), 17)  # wdExportFormatPDF
        finally:
            doc.Close(0)
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()

# ---------------------------------------------------------------- build
def build(lang, data, titles, profile, cum, exam_total, out_path):
    t, pr = T[lang], profile[lang]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(0.5)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    lang_el = OxmlElement('w:lang'); lang_el.set(qn('w:val'), 'es-MX' if lang == 'es' else 'en-US'); st.element.rPr.append(lang_el)
    zoom = doc.settings.element.find(qn('w:zoom'))
    if zoom is not None and zoom.get(qn('w:percent')) is None:
        zoom.set(qn('w:percent'), '100')
    doc.core_properties.title = f"{profile['name']} – {t['title']}"
    doc.core_properties.author = profile['name']

    para(doc, profile['name'], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, t['title'], italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, profile['contact'].replace(' | ', '  |  '), size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

    printed = fmt_date(data['transcript_printed'], lang)
    para(doc, pr['university'], bold=True, before=4)
    label_line(doc, t['program'], pr['program'])
    label_line(doc, t['majors'], '; '.join(pr['majors']))
    label_line(doc, t['minors'], '; '.join(pr['minors']))
    label_line(doc, t['grad'], pr['expected_graduation'])
    label_line(doc, t['gpa'], gpa(cum['points'], cum['gpa_hours']))
    label_line(doc, t['hours'], str(cum['earned']))
    para(doc, t['asof'].format(date=printed), italic=True, size=10, before=2)

    exams = [b for b in data['blocks'] if b['kind'] == 'exam']
    if exams:
        heading(doc, t['exam'])
        groups = []  # merge adjacent blocks that share a display label (e.g. Proficiency into AP)
        for b in exams:
            label = t['exam_kinds'].get(b['source'], b['source'])
            if groups and groups[-1][0] == label:
                groups[-1][1].extend(b['courses'])
            else:
                groups.append((label, list(b['courses'])))
        for label, courses in groups:
            para(doc, label, bold=True, before=2, after=1, keep=True)
            course_table(doc, courses, titles, lang)

    for b in (b for b in data['blocks'] if b['kind'] == 'term'):
        ip = all(not c['grade'] for c in b['courses'])
        heading(doc, f"{t['seasons'][b['season']]} {b['year']}")
        course_table(doc, b['courses'], titles, lang, in_progress=ip)
        cr = sum(c['credits'] for c in b['courses'])
        if ip:
            line = t['term_line_ip'].format(cr=cr)
        elif b['_term']['gpa_hours']:
            line = t['term_line'].format(cr=cr, gpa=gpa(b['_term']['points'], b['_term']['gpa_hours']),
                                         cum=gpa(b['_cum']['points'], b['_cum']['gpa_hours']))
        else:
            line = t['term_line_nogpa'].format(cr=cr)
        para(doc, line, italic=True, size=10, before=2)

    para(doc, t['key'], size=9, before=12, color='444444')
    para(doc, t['footer'].format(name=profile['name'], date=printed), size=9, before=4, color='444444')
    doc.save(out_path)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data', default='transcript-data.json')
    ap.add_argument('--docx-dir', default=os.path.join(HERE, 'output'))
    ap.add_argument('--pdf-dir', default=os.path.normpath(os.path.join(HERE, '..', '..', 'assets', 'pdfs', 'transcript')))
    ap.add_argument('--date', help='YYYYMMDD for the file names (default: the transcript print date)')
    ap.add_argument('--pdf', action='store_true', help='also export PDF with Word (private instance)')
    ap.add_argument('--strict', action='store_true', help='fail on unverified course titles')
    a = ap.parse_args()
    data, titles, profile = load(a.data), load('course-titles.json'), load('profile.json')
    cum, exam_total = check(data, titles, a.strict)          # TU's own numbers must match first
    adj_path = os.path.join(HERE, 'adjustments.json')
    if os.path.exists(adj_path):
        for n in apply_adjustments(data, load('adjustments.json')):
            print('ADJUSTMENT:', n)
        for c in (c for b in data['blocks'] for c in b['courses']):
            if title_for(c, titles) is None:
                sys.exit(f"STOPPED: no title for added course {c['code']}")
        cum = totals(data)                                    # display totals include the additions
    stamp = a.date or data['transcript_printed'].replace('-', '')
    os.makedirs(a.docx_dir, exist_ok=True)
    for lang in ('en', 'es'):
        name = f'{lang} Gideon Ong Transcript {stamp}'
        path = os.path.join(a.docx_dir, name + '.docx')
        build(lang, data, titles, profile, cum, exam_total, path)
        print('wrote', path)
        if a.pdf:
            os.makedirs(a.pdf_dir, exist_ok=True)
            pdf = os.path.join(a.pdf_dir, name + '.pdf')
            export_pdf(path, pdf)  # needs Microsoft Word (Windows)
            print('wrote', pdf)

if __name__ == '__main__':
    main()
